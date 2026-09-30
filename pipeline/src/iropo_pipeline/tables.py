"""Parse the FBI's NIBRS "Offense Type by Agency" tables.

Two layouts are handled:

* 2020-2023: one workbook per state (``North_Carolina_Offense_Type_by_Agency_2023.xlsx``). Agency
  type is only filled on the first row of each group, and some agencies are nested under group
  headers such as ``State Park Rangers:`` using cell indentation.
* 2024 onward: a single national workbook with a ``State/Territory`` column.

Only the columns IROPO needs are extracted: population, total offenses and animal cruelty.
"""

from __future__ import annotations

import io
import re
import zipfile
from collections.abc import Iterable, Iterator
from dataclasses import dataclass
from pathlib import Path

import openpyxl
import xlrd

from iropo_pipeline.states import lookup_state

AGENCY_TYPES = {
    "cities": "city",
    "city": "city",
    "metropolitan counties": "metro_county",
    "metropolitan county": "metro_county",
    "nonmetropolitan counties": "nonmetro_county",
    "nonmetropolitan county": "nonmetro_county",
    "universities and colleges": "university",
    "university or college": "university",
    "tribal": "tribal",
    "state police": "state_police",
    "other": "other",
    "other state agency": "other",
}

_FOOTNOTE_DIGITS = re.compile(r"(?<=[a-z])\d+\b")
# Footnote markers glued to agency names, e.g. "University of Oklahoma, Tulsa5".
_FOOTNOTE_SUFFIX = re.compile(r"(?<=[a-z)])\d$")


@dataclass(frozen=True)
class AgencyTableRow:
    year: int
    state: str
    agency_type: str
    name: str
    population: int | None
    total_offenses: int | None
    animal_cruelty: int | None


def normalize_header(value: object) -> str:
    """Normalize a column label: ``"Animal \\nCruelty"`` -> ``"animal cruelty"``.

    Footnote markers such as the ``1`` in ``Population1`` are removed.
    """
    if value is None:
        return ""
    text = str(value).lower().replace("\u2212", "-")
    text = _FOOTNOTE_DIGITS.sub("", text)
    text = re.sub(r"-\s+", "", text)  # re-join hyphenated line breaks ("Imper-\nsonation")
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return " ".join(text.split())


def normalize_agency_type(value: object) -> str | None:
    return AGENCY_TYPES.get(normalize_header(value))


def parse_int(value: object) -> int | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(round(value))
    text = str(value).strip().replace(",", "")
    if not text or text in {"-", "—", "N/A", "NA"}:
        return None
    try:
        return int(float(text))
    except ValueError:
        return None


def state_from_filename(filename: str) -> str | None:
    stem = Path(filename).stem
    match = re.match(r"(.+?)_Offense_Type_by_Agency", stem)
    if not match:
        return None
    state = lookup_state(match.group(1).replace("_", " "))
    return state.abbr if state else None


def state_from_title(rows: list[tuple]) -> str | None:
    """Resolve the state from a workbook's title rows (e.g. "PENNSYLVANIA")."""
    for row in rows[:3]:
        for cell in row[:2]:
            value = getattr(cell, "value", None)
            if isinstance(value, str) and (state := lookup_state(value)):
                return state.abbr
    return None


def _find_columns(rows: list[tuple]) -> tuple[int, dict[str, int]]:
    """Locate the header rows and map the columns we need. Returns (first data row, columns).

    Some years put every column label on one header row; others put offense names on a second
    row beneath grouped headings ("Crimes Against Society").
    """
    for index, row in enumerate(rows[:15]):
        labels = [normalize_header(cell.value) for cell in row]
        if "agency name" not in labels:
            continue
        two_rows = "animal cruelty" not in labels and index + 1 < len(rows)
        sub = [normalize_header(cell.value) for cell in rows[index + 1]] if two_rows else []
        columns: dict[str, int] = {}
        for col in range(max(len(labels), len(sub))):
            for label in (
                labels[col] if col < len(labels) else "",
                sub[col] if col < len(sub) else "",
            ):
                if label in {"state territory", "state", "federal state"}:
                    columns.setdefault("state", col)
                elif label == "agency type":
                    columns.setdefault("type", col)
                elif label == "agency name":
                    columns.setdefault("name", col)
                elif label == "population":
                    columns.setdefault("population", col)
                elif label == "total offenses":
                    columns.setdefault("total", col)
                elif label == "animal cruelty":
                    columns.setdefault("animal_cruelty", col)
        missing = {"type", "name", "total", "animal_cruelty"} - columns.keys()
        if missing:
            raise ValueError(f"Could not find columns {sorted(missing)} in table header")
        return (index + 2 if two_rows else index + 1), columns
    raise ValueError("No 'Agency Name' header row found")


@dataclass(frozen=True)
class SimpleCell:
    """Cell stand-in for legacy ``.xls`` files read with xlrd."""

    value: object
    indent: float = 0.0


def _cell(row: tuple, col: int | None):
    if col is None or col >= len(row):
        return None
    return row[col].value


def _indent(row: tuple, col: int) -> float:
    if col >= len(row):
        return 0.0
    cell = row[col]
    if isinstance(cell, SimpleCell):
        return cell.indent
    alignment = getattr(cell, "alignment", None)
    return float(getattr(alignment, "indent", 0) or 0)


def parse_worksheet_rows(
    rows: list[tuple], year: int, default_state: str | None
) -> Iterator[AgencyTableRow]:
    start, cols = _find_columns(rows)
    state = default_state
    agency_type: str | None = None
    group: str | None = None
    for row in rows[start:]:
        raw_name = _cell(row, cols["name"])
        raw_state = _cell(row, cols.get("state"))
        raw_type = _cell(row, cols["type"])
        if raw_state:
            resolved = lookup_state(str(raw_state))
            state = resolved.abbr if resolved else None
        if raw_type:
            agency_type = normalize_agency_type(raw_type)
        if raw_name is None or not str(raw_name).strip():
            continue  # blank line or footnote text in the first column
        name = _FOOTNOTE_SUFFIX.sub("", " ".join(str(raw_name).split()))
        total = parse_int(_cell(row, cols["total"]))
        cruelty = parse_int(_cell(row, cols["animal_cruelty"]))
        if name.endswith(":") and total is None and cruelty is None:
            group = name.rstrip(":").strip()
            continue
        if group and _indent(row, cols["name"]) > 0:
            name = f"{group}, {name}"
        else:
            group = None
        if state is None or agency_type is None:
            continue  # territories (e.g. Guam) and anything we cannot place
        yield AgencyTableRow(
            year=year,
            state=state,
            agency_type=agency_type,
            name=name,
            population=parse_int(_cell(row, cols.get("population"))),
            total_offenses=total,
            animal_cruelty=cruelty,
        )


def _is_agency_workbook(filename: str) -> bool:
    lower = Path(filename).name.lower()
    if not lower.endswith((".xlsx", ".xls")) or lower.startswith("~$"):
        return False
    if "__macosx" in filename.lower():
        return False
    if "federal" in lower and "tribal" not in lower:
        return False  # federal-only tables have no state jurisdictions
    # 2020-2023 archives also contain a national summary workbook without agency rows.
    return not lower.startswith("united_states_offense_type_by_agency")


def _xls_rows(data: bytes) -> list[tuple]:
    book = xlrd.open_workbook(file_contents=data, formatting_info=True)
    sheet = book.sheet_by_index(0)
    rows = []
    for r in range(sheet.nrows):
        cells = []
        for c in range(sheet.ncols):
            cell = sheet.cell(r, c)
            value = cell.value
            if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK) or value == "":
                value = None
            xf = book.xf_list[cell.xf_index] if cell.xf_index is not None else None
            cells.append(SimpleCell(value, float(xf.alignment.indent_level) if xf else 0.0))
        rows.append(tuple(cells))
    return rows


def _xlsx_rows(data: bytes) -> list[tuple]:
    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        return list(workbook.worksheets[0].iter_rows())
    finally:
        workbook.close()


def parse_agency_tables(zip_path: Path, year: int, log=print) -> list[AgencyTableRow]:
    rows: list[AgencyTableRow] = []
    with zipfile.ZipFile(zip_path) as archive:
        for info in archive.infolist():
            if not _is_agency_workbook(info.filename):
                continue
            data = archive.read(info)
            is_xls = info.filename.lower().endswith(".xls")
            sheet_rows = _xls_rows(data) if is_xls else _xlsx_rows(data)
            # File names are occasionally misspelled ("Pennyslvania"); fall back to the title.
            state = state_from_filename(info.filename) or state_from_title(sheet_rows)
            if state is None and "state" not in _find_columns(sheet_rows)[1]:
                log(f"  skipped {Path(info.filename).name} ({year}): not a U.S. state or DC")
                continue
            rows.extend(parse_worksheet_rows(sheet_rows, year, state))
    return rows


def dedupe(rows: Iterable[AgencyTableRow]) -> list[AgencyTableRow]:
    """Drop exact duplicates (the same agency listed twice in one state-year)."""
    seen: set[tuple] = set()
    unique = []
    for row in rows:
        key = (row.year, row.state, row.agency_type, row.name.lower())
        if key in seen:
            continue
        seen.add(key)
        unique.append(row)
    return unique
