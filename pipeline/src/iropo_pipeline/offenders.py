"""Parse the FBI's NIBRS offender demographic tables (age, sex, race, adult/juvenile)."""

from __future__ import annotations

import io
import re
import zipfile
from pathlib import Path

import openpyxl

from iropo_pipeline.tables import normalize_header, parse_int

DIMENSIONS = ("age_category", "age", "sex", "race")

AGE_KEYS = {
    "10 and under": "0-10",
    "66 and over": "66+",
    "unknown age": "unknown",
}
LABEL_KEYS = {
    "age_category": {"adult": "adult", "juvenile": "juvenile", "unknown age": "unknown"},
    "sex": {"male": "male", "female": "female", "unknown sex": "unknown"},
    "race": {
        "white": "white",
        "black or african american": "black",
        "american indian or alaska native": "american_indian_alaska_native",
        "asian": "asian",
        "native hawaiian or other pacific islander": "native_hawaiian_pacific_islander",
        "unknown race": "unknown",
    },
}
ROWS_OF_INTEREST = {"animal cruelty": "animal_cruelty", "total": "all_offenses"}


def dimension_for(filename: str) -> str | None:
    lower = Path(filename).name.lower()
    if "adult_and_juvenile" in lower:
        return "age_category"
    if "_age_by_" in lower:
        return "age"
    if "_sex_by_" in lower:
        return "sex"
    if "_race_by_" in lower:
        return "race"
    return None


def _label_key(dimension: str, label: str) -> str | None:
    if dimension == "age":
        if label in AGE_KEYS:
            return AGE_KEYS[label]
        match = re.fullmatch(r"(\d+)\s*(\d+)", label)  # "11 15" after normalization
        return f"{match.group(1)}-{match.group(2)}" if match else None
    return LABEL_KEYS[dimension].get(label)


def _label_columns(row: tuple, dimension: str) -> dict[int, str]:
    columns: dict[int, str] = {}
    for col, value in enumerate(row):
        key = _label_key(dimension, normalize_header(value)) if value is not None else None
        if key:
            columns[col] = key
    return columns


def parse_offender_sheet(rows: list[tuple], dimension: str) -> dict[str, dict[str, int]]:
    """Return ``{"animal_cruelty": {"total": n, <bucket>: n, ...}, "all_offenses": {...}}``."""
    header_index, columns = next(
        (
            (i, cols)
            for i, row in enumerate(rows[:12])
            if len(cols := _label_columns(row, dimension)) >= 2
        ),
        (None, {}),
    )
    if header_index is None:
        raise ValueError(f"No {dimension} columns found")
    result: dict[str, dict[str, int]] = {}
    for row in rows[header_index + 1 :]:
        if not row:
            continue
        name = ROWS_OF_INTEREST.get(normalize_header(row[0]))
        if not name:
            continue
        values = {"total": parse_int(row[1]) or 0}
        for col, key in columns.items():
            values[key] = parse_int(row[col] if col < len(row) else None) or 0
        result[name] = values
    missing = set(ROWS_OF_INTEREST.values()) - result.keys()
    if missing:
        raise ValueError(f"Rows {sorted(missing)} not found in {dimension} table")
    return result


def parse_offender_workbook(data: bytes | Path, dimension: str) -> dict[str, dict[str, int]]:
    source = io.BytesIO(data) if isinstance(data, bytes) else data
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    try:
        rows = [tuple(row) for row in workbook.worksheets[0].iter_rows(values_only=True)]
    finally:
        workbook.close()
    return parse_offender_sheet(rows, dimension)


def parse_offender_tables(zip_path: Path) -> dict[str, dict[str, dict[str, int]]]:
    """Parse every demographic workbook in a year's ``offenders.zip``, keyed by dimension."""
    tables: dict[str, dict[str, dict[str, int]]] = {}
    with zipfile.ZipFile(zip_path) as archive:
        for info in archive.infolist():
            dimension = dimension_for(info.filename)
            if dimension and info.filename.lower().endswith(".xlsx"):
                tables[dimension] = parse_offender_workbook(archive.read(info), dimension)
    return tables
